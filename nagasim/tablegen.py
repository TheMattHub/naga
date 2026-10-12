"""Models generated from tables (PHASE6 W7 Q5, ARCHITECTURE D128).

A planner already has the plant in spreadsheets: which stations exist, what
each product visits and in what order, how long each step takes. `generate()`
turns that directly into a runnable model, so the first draft of a model is a
table edit, not a canvas session.

    generate(source) -> layout

*source* is either a folder holding `Stations.csv`/`Products.csv`/
`Routings.csv` (and, optionally, `Orders.csv`/`Resources.csv`/`Schedules.csv`),
or one `.xlsx` workbook with a sheet per table. Every sheet/file is optional
except `Stations` and `Routings`, plus either `Products` or `Orders` (one
release mechanism is required). The path is the CALLER's, exactly like every
other CLI/API path in this codebase (netguard rule, ARCHITECTURE section 10.3)
-- `generate` opens whatever path it is given with no permission check of its
own.

**The model built.** One CREATE per `Products` row (or one Orders-mode CREATE
when an `Orders` table exists), one STATION + one PROCESS per `Stations` row,
one shared ROUTE dispatching by sequence (ARCHITECTURE 5.16/5.18), one
DISPOSE, and one `sequences` entry per distinct `product` in `Routings`,
ordered by `step`. A station with `x`/`y` also gets a `network` node (named by
its `node` cell, default the station name) and its `node` param set
(ARCHITECTURE 5.21), and every pair of consecutive routing stations on nodes
gets one two-way segment, so the ROUTE can always travel. Canvas positions come from the existing
auto-layout (`core/layout_tools.auto_layout`), not from `x`/`y` -- those are
LAYOUT (floor-plan) coordinates, a different space than the schematic canvas.

**D128, engine vs product (ARCHITECTURE section 10, O8).** This module sits
with the ENGINE: everything it imports (`core.document`, `core.layout_tools`,
`core.distributions`, `core.orders`) is already engine-side, and openpyxl is
already an engine dependency (`packages/engine/pyproject.toml`). Unlike
`modelcheck.py` (W7 Q3), tablegen calls nothing product-only, so the plan's
engine default holds with no exception needed; `packages/engine/build_wheel.py`'s
ENGINE list carries this file.

**The process-time ceiling (amended 2026-09-26 by Matteo's review, item 3).**
One PROCESS per station (ARCHITECTURE D51/D59) means a `Routings` row's
`process_time` cannot literally change that block's parameters per product --
only an EXPRESSION reading an entity attribute can. So every routing step
stamps `entity.process_time` (the row's own value, or the station's declared
distribution's MEAN when the cell is blank -- exact for every registered
distribution via `core/distributions.distribution_mean`) and a precomputed
`entity.time_scale` factor (`process_time / station_mean`; 1 when the cell is
blank). The PROCESS block's `dist_params` come from
`core/distributions.scale_distribution_params`: for a distribution that
declares `scale_params` (a location-scale family -- Constant, Uniform,
Exponential, Normal, Lognormal, Triangular, Gamma, Weibull, Erlang, Rayleigh,
LogLogistic, JohnsonSU, PearsonTypeV/VI; see the registry comments in
`core/distributions.py` for the full, decided list), every scale/location
parameter becomes an expression of `entity.time_scale` and every shape
parameter stays the station's own literal value -- so the FAMILY is always
kept, and its mean at that visit is always the `Routings` row's own value.
A distribution with no `scale_params` (an empirical table, CustomExpression,
FixedSchedule, or a bounded family whose bounds are not themselves scale
params, e.g. BoundedNormal/Beta/JohnsonSB/the discrete counting families)
keeps its own literal `param1`/`param2` for every product; a `Routings` row
that tries to override `process_time` at such a station is a `ValueError`
naming the station, the product and the distribution -- there is no silent
"override has no effect" any more. Two more stated ceilings, both structural,
not just of this implementation: a `Routings` sequence is linear (no
branching, no rework loop -- `core/sequences.py`'s own model), and there is no
batching (one entity per routing step, matching the "no batching" rule the
rest of the table contract assumes).

CLI: `naga generate TABLES -o model.json`, `naga generate --template
out.xlsx`. Editor: File > New from ta&bles... (`product/ui/main_window.py`).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import sys
from pathlib import Path

from nagasim.core.distributions import (distribution_mean, distribution_scale_params,
                                        get_distribution_params, scale_distribution_params)
from nagasim.core.document import ModelDocument
from nagasim.core.layout_tools import auto_layout
from nagasim.core.orders import CSV_COLUMNS as ORDER_CSV_COLUMNS
from nagasim.core.orders import ORDERS_MODE, orders_from_csv

__all__ = ["generate", "write_template", "main",
           "STATIONS_COLUMNS", "PRODUCTS_COLUMNS", "ROUTINGS_COLUMNS",
           "RESOURCES_COLUMNS", "SCHEDULES_COLUMNS"]

# The column contract (independent oracle tests/test_w7_q5_oracle.py fixes
# these three -- adopted verbatim). `start_time` is the oracle's own addition
# beyond the plan text's four guessed Products columns: without a per-product
# arrival offset, two products released at t=0 race for a shared station in
# whatever order CREATE blocks happen to be scheduled in the same instant.
STATIONS_COLUMNS = ["name", "resource", "capacity", "distribution", "param1", "param2",
                    "node", "x", "y"]
PRODUCTS_COLUMNS = ["entity_type", "arrival_distribution", "arrival_param1", "max_arrivals",
                    "start_time"]
ROUTINGS_COLUMNS = ["product", "step", "station", "process_time"]
# Resources/Schedules (plan text: "optional Resources, Schedules"): a thin
# passthrough, not a friendlier per-column shape -- `points`/a schedule's own
# shape is one JSON cell. Add a richer form if a real project needs one.
RESOURCES_COLUMNS = ["name", "capacity", "discipline", "cost_busy", "cost_idle", "cost_per_use",
                     "cost_time_unit", "capacity_schedule"]
SCHEDULES_COLUMNS = ["name", "kind", "points", "repeat"]

_TABLE_NAMES = ("Stations", "Products", "Routings", "Orders", "Resources", "Schedules")


# ---------------------------------------------------------------- reading

def _str(value) -> str:
    return "" if value is None else str(value).strip()


def _num(value):
    """*value* as a float, or None for blank/unparseable -- a bool is never a
    number here (Python's bool is an int subclass, which would silently turn
    a stray True/False cell into 1.0/0.0)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _csv_table(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        header = list(reader.fieldnames or [])
        rows = [dict(row) for row in reader]
    return header, rows


def _xlsx_table(ws) -> tuple[list[str], list[dict]]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return [], []
    header = [_str(h) for h in values[0]]
    rows = []
    for raw in values[1:]:
        if all(v is None for v in raw):
            continue
        rows.append({header[i]: raw[i] for i in range(len(header)) if i < len(raw)})
    return header, rows


def _read_tables(source) -> dict[str, tuple[list[str], list[dict]]]:
    path = Path(source)
    if path.is_dir():
        tables = {}
        for name in _TABLE_NAMES:
            csv_path = path / f"{name}.csv"
            if csv_path.is_file():
                tables[name] = _csv_table(csv_path)
        return tables
    if path.is_file():
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        return {name: _xlsx_table(wb[name]) for name in _TABLE_NAMES if name in wb.sheetnames}
    raise ValueError(f"tablegen: source not found: {source}")


def _check_columns(sheet: str, header: list[str], required: tuple[str, ...]) -> None:
    missing = [c for c in required if c not in header]
    if missing:
        raise ValueError(f"{sheet}: missing required column(s): {', '.join(missing)}")


# ---------------------------------------------------------------- stations

def _add_stations(doc: ModelDocument, table) -> dict[str, dict]:
    header, rows = table
    _check_columns("Stations", header, ("name", "capacity", "distribution", "param1"))
    station_info: dict[str, dict] = {}
    nodes: dict[str, dict] = {}
    for row in rows:
        name = _str(row.get("name"))
        if not name:
            raise ValueError("Stations: a row is missing 'name'")
        if name in station_info:
            raise ValueError(f"Stations: duplicate station name '{name}'")
        capacity = _num(row.get("capacity"))
        if capacity is None:
            raise ValueError(f"Stations: station '{name}' is missing 'capacity'")
        resource = _str(row.get("resource")) or name
        dist = _str(row.get("distribution")) or "Constant"
        param_names = get_distribution_params(dist)
        if not param_names:
            raise ValueError(f"Stations: station '{name}' names an unknown distribution '{dist}'")

        dist_params: dict[str, float] = {}
        p1 = _num(row.get("param1"))
        if p1 is not None:
            dist_params[param_names[0]] = p1
        p2 = _num(row.get("param2"))
        if p2 is not None and len(param_names) > 1:
            dist_params[param_names[1]] = p2

        default_time = distribution_mean(dist, dist_params)
        if default_time is None:
            raise ValueError(
                f"Stations: station '{name}': '{dist}' with {dist_params} has no finite mean; "
                "give every Routings row through this station its own process_time")

        x, y = _num(row.get("x")), _num(row.get("y"))
        node_name = None
        if x is not None and y is not None:
            node_name = _str(row.get("node")) or name
            if node_name not in nodes:      # stations sharing a node: the first row places it
                nodes[node_name] = {"name": node_name, "x": x, "y": y}

        block = doc.add_block("STATION", (0.0, 0.0), name=name)
        if node_name:
            params = dict(doc.blocks[block].params)
            params["node"] = node_name
            doc.set_params(block, params)

        station_info[name] = {"resource": resource, "capacity": int(capacity), "distribution": dist,
                              "dist_params": dist_params, "default_time": default_time,
                              "node": node_name}

    if nodes:
        doc.set_network({"nodes": list(nodes.values()), "segments": []})

    resources: dict[str, dict] = {}
    owner: dict[str, str] = {}
    for name, info in station_info.items():
        entry = resources.setdefault(info["resource"], {"name": info["resource"],
                                                         "capacity": info["capacity"],
                                                         "discipline": "FIFO"})
        owner.setdefault(info["resource"], name)
        if entry["capacity"] != info["capacity"]:
            raise ValueError(
                f"Stations: '{owner[info['resource']]}' and '{name}' share resource "
                f"'{info['resource']}' with different capacities "
                f"({entry['capacity']} and {info['capacity']})")
    doc.set_resources(list(resources.values()))
    return station_info


def _segments(sequences: list[dict], station_info: dict) -> list[dict]:
    """One two-way segment per pair of consecutive stations of every routing
    that both sit on a node, deduplicated in either direction. No `length`:
    the network takes the straight-line distance between the two nodes'
    coordinates (core/network.py), so a routing's travel is always reachable."""
    out, seen = [], set()
    for sequence in sequences:
        path = [station_info[s["station"]]["node"] for s in sequence["steps"]]
        for a, b in zip(path, path[1:]):
            if a and b and a != b and frozenset((a, b)) not in seen:
                seen.add(frozenset((a, b)))
                out.append({"from": a, "to": b, "two_way": True})
    return out


def _add_station_process(doc: ModelDocument, name: str, info: dict, route_name: str) -> None:
    process_name = doc.add_block("PROCESS", (0.0, 0.0), name=f"process_{name.lower()}")
    # Any distribution with `scale_params` gets every scale/location parameter
    # rewritten as an expression of the per-visit `entity.time_scale` factor
    # (module docstring); one without keeps its own literal dist_params, since
    # `_build_sequences` already refused a Routings override at such a station.
    scaled = scale_distribution_params(info["distribution"], info["dist_params"], "entity.time_scale")
    dist_params = scaled if scaled is not None else dict(info["dist_params"])
    params = dict(doc.blocks[process_name].params)
    params.update({"action": "Seize Delay Release", "distribution": info["distribution"],
                  "dist_params": dist_params, "resource_name": info["resource"]})
    doc.set_params(process_name, params)
    doc.add_link(name, process_name)
    doc.add_link(process_name, route_name)


# ---------------------------------------------------------------- routings/sequences

def _build_sequences(table, station_info: dict) -> list[dict]:
    header, rows = table
    _check_columns("Routings", header, ("product", "step", "station"))
    by_product: dict[str, list[tuple[float, str, float | None]]] = {}
    for row in rows:
        product = _str(row.get("product"))
        station = _str(row.get("station"))
        step = _num(row.get("step"))
        if not product or not station or step is None:
            raise ValueError("Routings: every row needs 'product', 'step' and 'station'")
        if station not in station_info:
            raise ValueError(f"Routings: unknown station '{station}'")
        process_time = _num(row.get("process_time"))
        distribution = station_info[station]["distribution"]
        if process_time is not None and distribution_scale_params(distribution) is None:
            raise ValueError(
                f"Routings: product '{product}' overrides process_time at station '{station}', but "
                f"'{distribution}' has no rescalable parameters (an empirical table, "
                f"CustomExpression, FixedSchedule, and a bounded family whose bounds are not "
                f"themselves scale params keep one literal shape for every product); drop the "
                f"override or give '{station}' a scale-able distribution")
        by_product.setdefault(product, []).append((step, station, process_time))

    sequences = []
    for product, steps in by_product.items():
        steps.sort(key=lambda s: s[0])
        seq_steps = []
        for _, station, process_time in steps:
            station_mean = station_info[station]["default_time"]
            target = process_time if process_time is not None else station_mean
            factor = target / station_mean if station_mean else 1.0
            seq_steps.append({"station": station,
                              "values": {"process_time": target, "time_scale": factor}})
        sequences.append({"name": product, "steps": seq_steps})
    return sequences


# ---------------------------------------------------------------- products / orders

def _add_product_creates(doc: ModelDocument, table, route_name: str, routed: set) -> None:
    header, rows = table
    _check_columns("Products", header,
                   ("entity_type", "arrival_distribution", "arrival_param1", "max_arrivals"))
    for row in rows:
        e_type = _str(row.get("entity_type"))
        if not e_type:
            raise ValueError("Products: a row is missing 'entity_type'")
        if e_type not in routed:
            raise ValueError(f"Products: product '{e_type}' has no Routings rows")
        dist = _str(row.get("arrival_distribution")) or "Constant"
        param_names = get_distribution_params(dist)
        if not param_names:
            raise ValueError(f"Products: entity '{e_type}' names an unknown distribution '{dist}'")
        dist_params: dict[str, float] = {}
        p1 = _num(row.get("arrival_param1"))
        if p1 is not None:
            dist_params[param_names[0]] = p1
        max_arrivals = _num(row.get("max_arrivals"))
        start_time = _num(row.get("start_time"))

        name = doc.add_block("CREATE", (0.0, 0.0), name=f"create_{e_type.lower()}")
        params = dict(doc.blocks[name].params)
        params.update({
            "e_type": e_type, "arrival_mode": "Interarrival",
            "distribution": dist, "dist_params": dist_params,
            "entities_per_arrival": 1,
            # CREATE's start_time/max_arrivals are typed int (block_spec.py),
            # not float -- a fractional start_time is outside this table's
            # contract, same ceiling as the engine's own.
            "max_arrivals": int(max_arrivals) if max_arrivals is not None else 0,
            "start_time": int(start_time) if start_time is not None else 0,
            "sequence": e_type,
        })
        doc.set_params(name, params)
        doc.add_link(name, route_name)


def _cell(value):
    return "" if value is None else value


def _orders_section(rows: list[dict]) -> list[dict]:
    """Rows (from either source shape) -> the model's "orders" section, by
    round-tripping through `core/orders.py`'s own CSV reader instead of
    duplicating its parsing rules."""
    if not rows:
        return []
    extras = [k for k in rows[0].keys() if k and k not in ORDER_CSV_COLUMNS]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=[*ORDER_CSV_COLUMNS, *extras])
    writer.writeheader()
    for row in rows:
        writer.writerow({col: _cell(row.get(col)) for col in (*ORDER_CSV_COLUMNS, *extras)})
    return orders_from_csv(buf.getvalue())


def _add_orders_create(doc: ModelDocument, table, route_name: str) -> None:
    header, rows = table
    _check_columns("Orders", header, ("id",))
    orders = _orders_section(rows)
    name = doc.add_block("CREATE", (0.0, 0.0), name="create_orders")
    params = dict(doc.blocks[name].params)
    params["arrival_mode"] = ORDERS_MODE
    doc.set_params(name, params)
    doc.add_link(name, route_name)
    doc.set_orders(orders)


# ---------------------------------------------------------------- optional Resources/Schedules

def _apply_resources_sheet(doc: ModelDocument, table) -> None:
    header, rows = table
    _check_columns("Resources", header, ("name",))
    by_name = {r["name"]: dict(r) for r in doc.resources}
    for row in rows:
        name = _str(row.get("name"))
        if not name:
            continue
        entry = by_name.setdefault(name, {"name": name, "capacity": 1, "discipline": "FIFO"})
        capacity = _num(row.get("capacity"))
        if capacity is not None:
            entry["capacity"] = capacity
        discipline = _str(row.get("discipline"))
        if discipline:
            entry["discipline"] = discipline
        for key in ("cost_busy", "cost_idle", "cost_per_use"):
            value = _num(row.get(key))
            if value is not None:
                entry[key] = value
        cost_time_unit = _str(row.get("cost_time_unit"))
        if cost_time_unit:
            entry["cost_time_unit"] = cost_time_unit
        schedule_name = _str(row.get("capacity_schedule"))
        if schedule_name:
            entry["capacity_schedule"] = schedule_name
    doc.set_resources(list(by_name.values()))


def _schedules_from_rows(table) -> list[dict]:
    header, rows = table
    _check_columns("Schedules", header, ("name", "points"))
    schedules = []
    for row in rows:
        name = _str(row.get("name"))
        if not name:
            continue
        raw_points = row.get("points")
        try:
            if isinstance(raw_points, str) and raw_points.strip():
                points = json.loads(raw_points)
            else:
                points = raw_points or []
        except json.JSONDecodeError as exc:
            raise ValueError(f"Schedules: '{name}': 'points' is not valid JSON: {exc}") from None
        entry = {"name": name, "kind": _str(row.get("kind")) or "capacity", "points": points}
        repeat = _num(row.get("repeat"))
        if repeat is not None:
            entry["repeat"] = repeat
        schedules.append(entry)
    return schedules


# ---------------------------------------------------------------- generate

def generate(source) -> dict:
    """A full model layout from *source* (an .xlsx workbook path, or a folder
    of CSVs). Raises ValueError naming the problem: a missing required
    sheet/column, or a Routings row naming a station that isn't in Stations."""
    tables = _read_tables(source)

    if "Stations" not in tables:
        raise ValueError("tablegen: missing required table 'Stations'")
    if "Routings" not in tables:
        raise ValueError("tablegen: missing required table 'Routings'")
    has_orders = "Orders" in tables
    if not has_orders and "Products" not in tables:
        raise ValueError("tablegen: missing required table 'Products' (or an 'Orders' table)")

    doc = ModelDocument()
    station_info = _add_stations(doc, tables["Stations"])
    sequences = _build_sequences(tables["Routings"], station_info)
    doc.set_sequences(sequences)
    if doc.network:
        doc.set_network({**doc.network, "segments": _segments(sequences, station_info)})

    route_name = doc.add_block("ROUTE", (0.0, 0.0))
    dispose_name = doc.add_block("DISPOSE", (0.0, 0.0))
    doc.add_link(route_name, dispose_name)

    for name, info in station_info.items():
        _add_station_process(doc, name, info, route_name)

    if has_orders:
        _add_orders_create(doc, tables["Orders"], route_name)
    else:
        _add_product_creates(doc, tables["Products"], route_name,
                             {seq["name"] for seq in sequences})

    if "Resources" in tables:
        _apply_resources_sheet(doc, tables["Resources"])
    if "Schedules" in tables:
        doc.set_schedules(_schedules_from_rows(tables["Schedules"]))

    doc.move_blocks(auto_layout(doc))
    return doc.to_layout()


# ---------------------------------------------------------------- template / CLI

def write_template(path) -> None:
    """An empty .xlsx workbook, one sheet per recognised table, header row
    only -- what a planner fills in and hands back to `generate`."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, columns in (("Stations", STATIONS_COLUMNS), ("Products", PRODUCTS_COLUMNS),
                          ("Routings", ROUTINGS_COLUMNS), ("Orders", list(ORDER_CSV_COLUMNS)),
                          ("Resources", RESOURCES_COLUMNS), ("Schedules", SCHEDULES_COLUMNS)):
        ws = wb.create_sheet(name)
        ws.append(columns)
    wb.save(str(path))


def main(argv=None) -> int:
    """`naga generate TABLES -o model.json` and `naga generate --template
    out.xlsx`. Exit 0 on success, 2 on a bad source/table."""
    parser = argparse.ArgumentParser(
        prog="naga generate",
        description="Generate a model from Stations/Products/Routings tables "
                    "(an .xlsx workbook or a folder of CSVs).")
    parser.add_argument("source", nargs="?", help="an .xlsx workbook, or a folder of CSVs")
    parser.add_argument("-o", "--output", help="write the generated model JSON here (default: stdout)")
    parser.add_argument("--template", metavar="FILE",
                        help="write an empty .xlsx workbook with header rows and exit")
    args = parser.parse_args(argv)

    if args.template:
        try:
            write_template(args.template)
        except OSError as exc:
            print(f"naga generate: {exc}", file=sys.stderr)
            return 2
        print(f"wrote template: {args.template}")
        return 0

    if not args.source:
        parser.error("a source (.xlsx file or folder of CSVs) is required unless --template is given")

    try:
        layout = generate(args.source)
    except (ValueError, OSError) as exc:
        print(f"naga generate: {exc}", file=sys.stderr)
        return 2

    text = json.dumps(layout, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"wrote model: {args.output}")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
