"""The study's deliverable: the control room page, an Excel workbook, the current model and
RO-Crate metadata, written to `<study>/package/`. Built from the study folder alone."""
from __future__ import annotations

import json
import os
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from nagasim import study as study_mod, studypage
from nagasim.studychecks import goal_kpi
from nagasim.twin import write_json, write_text

_FORMATS = {".html": "text/html", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ".json": "application/json",
            ".md": "text/markdown"}


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:   # study text is data: "=..." must not become a formula when Excel opens it
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)


def workbook(folder: Path, study: dict, path: Path) -> None:
    """Sheets Goals, Versions (one column per KPI seen in any version), Checklist, Assumptions."""
    folder, path = Path(folder), Path(path)
    last = (study["versions"][-1].get("kpis") or {}) if study["versions"] else {}
    goals = [[g["text"], goal_kpi(spec), spec.get("low"), spec.get("high"), last.get(goal_kpi(spec)),
              g["status"]] for spec, g in zip(study["goals"], study_mod.goals_status(study))]
    kpis = list(dict.fromkeys(k for v in study["versions"] for k in (v.get("kpis") or {})))
    versions = []
    for v in study["versions"]:
        step = studypage._read(folder / "versions" / v["id"] / "step.json") or {}
        versions.append([v["id"], v.get("note", ""), v.get("verdict", "not_verified"), step.get("at", "")]
                        + [(v.get("kpis") or {}).get(k) for k in kpis])
    checks = [[c["title"], c["verdict"], c["why"]] for c in study_mod.checklist(study, folder)]
    assumptions = [[a.get("text") or a.get("statement") or "", a.get("value"), a.get("source") or ""]
                   for a in study["assumptions"]]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Goals", ["Name", "KPI", "Low", "High", "Last value", "Status"], goals)
    _sheet(wb, "Versions", ["Id", "Note", "Verdict", "Date"] + kpis, versions)
    _sheet(wb, "Checklist", ["Item", "Verdict", "Reason"], checks)
    _sheet(wb, "Assumptions", ["Statement", "Value", "Source"], assumptions)
    tmp = path.with_name(f"{path.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def ro_crate(study: dict, files: list[str]) -> dict:
    """RO-Crate 1.1 JSON-LD, hand-written: the descriptor, the root dataset, one File per name."""
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    graph = [
        {"@id": "ro-crate-metadata.json", "@type": "CreativeWork",
         "about": {"@id": "./"}, "conformsTo": {"@id": "https://w3id.org/ro/crate/1.1"}},
        {"@id": "./", "@type": "Dataset", "name": study["question"], "description": study.get("decision", ""),
         "datePublished": now, "dateModified": now,
         "hasPart": [{"@id": f} for f in files]},
    ]
    graph += [{"@id": f, "@type": "File", "name": f,
               "encodingFormat": _FORMATS.get(Path(f).suffix, "application/octet-stream")} for f in files]
    return {"@context": "https://w3id.org/ro/crate/1.1/context", "@graph": graph}


def package(folder: Path, study: dict) -> list[str]:
    """Write the package; returns its paths (four, plus any write-up or slides) relative to the study folder."""
    folder = Path(folder)
    if not study["versions"]:
        raise study_mod.StudyError("E_STUDY_NO_VERSION", "the study has no model version yet",
                                   "Build a version with study_step first.")
    out = folder / "package"
    out.mkdir(exist_ok=True)
    write_text(out / "control-room.html", studypage.render(folder, study))
    workbook(folder, study, out / "study.xlsx")
    model = folder / "versions" / study["versions"][-1]["id"] / "model.json"
    tmp = out / f"model.json.{uuid.uuid4().hex[:8]}.tmp"
    try:
        shutil.copyfile(model, tmp)
        os.replace(tmp, out / "model.json")
    finally:
        tmp.unlink(missing_ok=True)
    names = ["control-room.html", "study.xlsx", "model.json"]
    # Written by study_writeup and study_slides, which record the version they were built from in
    # outputs.json; listed only when that is still the last version (an older one would mislead).
    built = studypage._read(out / "outputs.json") or {}
    last = study["versions"][-1]["id"]
    names += [n for key, files in (("writeup", ("writeup.html", "writeup.md")), ("slides", ("slides.html",)))
              if built.get(key) == last for n in files if (out / n).is_file()]
    write_json(out / "ro-crate-metadata.json", ro_crate(study, names))
    return [f"package/{n}" for n in names + ["ro-crate-metadata.json"]]
