"""The study package: workbook, RO-Crate, the four files (nagasim.studypackage)."""
import json

import openpyxl
import pytest

from nagasim import study, studypackage

_OK = [{"name": "c", "verdict": "passed", "why": "", "kind": "law"}]


def _layout(*names):
    return {"blocks": [{"name": n, "type": "Delay"} for n in names], "links": []}


def _study(tmp_path):
    folder = tmp_path / "s"
    st = study.create(folder, "Should we add a crane?")
    st["goals"] = [{"id": "g1", "text": "Throughput", "kpi": "A", "low": 5}]
    st["assumptions"] = [{"id": "a1", "text": "Arrivals are Poisson", "source": "log"}]
    study.record_step(folder, st, _layout("a"), {"system": {"A": {"mean": 7.0}}}, _OK, "first", None)
    study.record_step(folder, st, _layout("a", "b"), {"system": {"A": {"mean": 8.0}, "B": {"mean": 1.0}}},
                      _OK, "second", None)
    return folder, st


def test_workbook_sheets_and_kpi_columns(tmp_path):
    folder, st = _study(tmp_path)
    studypackage.workbook(folder, st, tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "w.xlsx")
    assert wb.sheetnames == ["Goals", "Versions", "Checklist", "Assumptions"]
    rows = list(wb["Versions"].values)
    assert rows[0][-2:] == ("A", "B") and rows[1][-1] is None and rows[2][-1] == 1.0
    assert list(wb["Goals"].values)[1][4:] == (8.0, "met")


def test_ro_crate_lists_files():
    crate = studypackage.ro_crate({"question": "q", "decision": "d"}, ["a.html", "b.xlsx"])
    ids = {e["@id"]: e for e in crate["@graph"]}
    assert crate["@context"] == "https://w3id.org/ro/crate/1.1/context"
    assert ids["./"]["name"] == "q" and ids["a.html"]["encodingFormat"] == "text/html"
    assert ids["b.xlsx"]["@type"] == "File"


def test_package_writes_four_files(tmp_path):
    folder, st = _study(tmp_path)
    names = studypackage.package(folder, st)
    assert len(names) == 4 and all((folder / n).is_file() for n in names)
    assert json.loads((folder / "package" / "model.json").read_text()) == _layout("a", "b")


def test_package_without_versions_is_a_coded_error(tmp_path):
    folder = tmp_path / "s"
    st = study.create(folder, "q")
    with pytest.raises(study.StudyError) as e:
        studypackage.package(folder, st)
    assert e.value.code == "E_STUDY_NO_VERSION"


def test_workbook_keeps_formula_like_text_as_text(tmp_path):
    from openpyxl import Workbook, load_workbook
    from nagasim.studypackage import _sheet
    wb = Workbook()
    _sheet(wb, "S", ["a"], [["=HYPERLINK(\"x\")"]])
    wb.save(tmp_path / "w.xlsx")
    cell = load_workbook(tmp_path / "w.xlsx")["S"]["A2"]
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("x")'


def test_package_lists_writeup_and_slides_when_they_exist(tmp_path):
    folder, st = _study(tmp_path)
    assert not any("writeup" in f or "slides" in f for f in studypackage.package(folder, st))
    for n in ("writeup.html", "writeup.md", "slides.html"):
        (folder / "package" / n).write_text("x", encoding="utf-8")
    outputs = folder / "package" / "outputs.json"
    outputs.write_text(json.dumps({"writeup": "v001", "slides": "v001"}), encoding="utf-8")
    assert not any("writeup" in f or "slides" in f for f in studypackage.package(folder, st))   # built from v001, last is v002
    outputs.write_text(json.dumps({"writeup": "v002", "slides": "v002"}), encoding="utf-8")
    files = studypackage.package(folder, st)
    assert {"package/writeup.html", "package/writeup.md", "package/slides.html"} <= set(files)
    crate = json.loads((folder / "package" / "ro-crate-metadata.json").read_text(encoding="utf-8"))
    ids = {e["@id"] for e in crate["@graph"]}
    assert {"writeup.html", "writeup.md", "slides.html"} <= ids


def test_page_on_request_lists_new_tools_and_sensitivity_table(tmp_path):
    from nagasim import studypage
    folder, st = _study(tmp_path)
    page = studypage.render(folder, st)
    assert "planned" not in page and "study_sensitivity" in page and 'id="sensitivity"' not in page
    vdir = folder / "versions" / st["versions"][-1]["id"]
    (vdir / "sensitivity.json").write_text(json.dumps({
        "charts": {"A": "<svg onload=evil></svg>"}, "changes_answer": ["x"],
        "factors": [{"path": "x", "label": "Arrival rate", "low": 1, "high": 2,
                     "kpis": {"A": {"swing": 3.5, "flips": True}}}]}), encoding="utf-8")
    page = studypage.render(folder, st)
    assert 'id="sensitivity"' in page and "Arrival rate" in page and "evil" not in page
